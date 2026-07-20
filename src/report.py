from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .config import Settings
from .parsers import ExtractedData

logger = logging.getLogger(__name__)

COLUNAS = [
    "Arquivo Original",
    "Nome Novo",
    "Data Extraída",
    "Recebedor Extraído",
    "Valor Extraído",
    "Layout Detectado",
    "Status",
    "Timestamp",
]

_pendentes: list[tuple] = []


def registrar(
    cfg: Settings,
    arquivo_original: str,
    nome_novo: str,
    dados: ExtractedData,
    status: str,
) -> None:
    linha = (
        arquivo_original,
        nome_novo,
        dados.data_pagamento or "",
        dados.nome_recebedor or "",
        dados.valor or "",
        dados.layout,
        status,
        datetime.now().isoformat(timespec="seconds"),
    )
    _pendentes.append(linha)
    _tentar_gravar_pendentes(cfg.caminho_relatorio)


def _tentar_gravar_pendentes(caminho_relatorio: Path) -> None:
    if not _pendentes:
        return
    try:
        wb = _abrir_ou_criar(caminho_relatorio)
        ws = wb.active
        for linha in _pendentes:
            ws.append(linha)
        caminho_relatorio.parent.mkdir(parents=True, exist_ok=True)
        wb.save(caminho_relatorio)
        _pendentes.clear()
    except PermissionError:
        logger.warning(
            "Não foi possível salvar %s (arquivo aberto em outro programa?). "
            "%d linha(s) pendente(s), será tentado novamente no próximo ciclo.",
            caminho_relatorio,
            len(_pendentes),
        )


def _abrir_ou_criar(caminho_relatorio: Path) -> Workbook:
    if caminho_relatorio.exists():
        return load_workbook(caminho_relatorio)
    wb = Workbook()
    ws = wb.active
    ws.append(COLUNAS)
    return wb
